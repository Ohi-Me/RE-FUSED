"""Parse the ERPC block-wise deviation settlement workbooks into one table for the Eastern-Region State buyers.

Input:  data/raw/erpc_dsm/*.xlsx (a20_erpc_dsm_blockwise.py), or the files given on the command line
Output: data/interim/erpc_dsm_blocks.parquet with one row per entity and 15-minute block:
        entity (panel code), date, block, freq_hz, actual_mwh, sched_mwh (drawal positive), dev_mwh (buyer sign:
        positive = over-drawal), payable_rs, receivable_rs, nr_rs_mwh, source file.
Where a revised account and the original cover the same week, the revised one is kept.
"""
import glob
import os
import sys

import numpy as np
import openpyxl
import pandas as pd

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
RAW = os.path.join(ROOT, "data", "raw", "erpc_dsm")
OUT = os.path.join(ROOT, "data", "interim", "erpc_dsm_blocks.parquet")
SHEETS = {"BSPHCL": "BR", "JBVNL": "JH", "GRIDCO": "OD", "WBSETCL": "WB", "DVC": "DVC", "SIKKIM": "SK"}


def parse(path):
    try:
        wb = openpyxl.load_workbook(open(path, "rb"), read_only=True, data_only=True)
    except Exception as e:                                    # not a workbook of this kind
        print("skipped", os.path.basename(path), type(e).__name__)
        return None
    out = []
    for sheet, code in SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = []
        for r in wb[sheet].iter_rows(min_row=5, values_only=True):
            if r[2] is None or not isinstance(r[2], (int, float)) or r[0] is None:
                continue
            rows.append(r[:12])
        if not rows:
            continue
        d = pd.DataFrame(rows, columns=["date", "time", "block", "freq_hz", "actual", "sched", "sras", "dev", "dev_pct",
                                        "payable_rs", "receivable_rs", "nr_p_kwh"])
        d["entity"] = code
        out.append(d)
    if not out:
        return None
    d = pd.concat(out)
    d["date"] = pd.to_datetime(d.date)
    d["block"] = d.block.astype(int)
    # drawal is negative in the accounts; buyer sign: sched and actual as drawal, deviation positive = over-drawal
    d["sched_mwh"] = -d.sched.astype(float)
    d["actual_mwh"] = -d.actual.astype(float)
    d["dev_mwh"] = -d.dev.astype(float)
    d["nr_rs_mwh"] = 10.0 * d.nr_p_kwh.astype(float)
    d["source"] = os.path.basename(path)
    d["revised"] = os.path.basename(path).startswith("Revised")
    return d[["entity", "date", "block", "freq_hz", "actual_mwh", "sched_mwh", "dev_mwh", "payable_rs",
              "receivable_rs", "nr_rs_mwh", "source", "revised"]]


def main():
    files = sys.argv[1:] or sorted(glob.glob(os.path.join(RAW, "*.xlsx")) + glob.glob(os.path.join(RAW, "*.bin")))
    parts = [p for p in (parse(f) for f in files) if p is not None]
    D = pd.concat(parts)
    D = D.sort_values("revised").drop_duplicates(["entity", "date", "block"], keep="last")
    D = D.sort_values(["entity", "date", "block"]).reset_index(drop=True)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    D.to_parquet(OUT)
    print(len(files), "files,", len(D), "rows,", D.date.min().date(), "to", D.date.max().date())
    print(D.groupby("entity").date.agg(["min", "max", "nunique"]))


if __name__ == "__main__":
    main()
